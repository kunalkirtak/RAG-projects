"""
document_loader.py
-------------------
Loads raw documents (PDF, DOCX, TXT) from disk, extracts text + metadata,
and cleans the extracted text so it is ready for chunking.
"""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

from src.logger import get_logger
from src.utils import clean_text, generate_doc_id

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class Document:
    """A single loaded source document."""

    doc_id: str
    source: str
    text: str
    metadata: Dict[str, Any] = field(default_factory=dict)


class DocumentLoadError(Exception):
    """Raised when a document cannot be parsed."""


def _load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise DocumentLoadError("pypdf is required to load PDF files") from exc

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _load_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentLoadError("python-docx is required to load DOCX files") from exc

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


_LOADERS = {
    ".txt": _load_txt,
    ".md": _load_txt,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
}


def load_document(path: str | Path) -> Document:
    """Load a single document from disk into a `Document` object.

    Args:
        path: Path to a .pdf, .docx, .txt, or .md file.

    Returns:
        A `Document` with cleaned text and extracted metadata.

    Raises:
        DocumentLoadError: If the file type is unsupported or parsing fails.
    """
    path = Path(path)
    ext = path.suffix.lower()

    if ext not in _LOADERS:
        raise DocumentLoadError(
            f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    if not path.exists():
        raise DocumentLoadError(f"File not found: {path}")

    try:
        raw_text = _LOADERS[ext](path)
    except DocumentLoadError:
        raise
    except Exception as exc:  # noqa: BLE001 - surface as a DocumentLoadError
        raise DocumentLoadError(f"Failed to parse {path.name}: {exc}") from exc

    text = clean_text(raw_text)
    stat = path.stat()
    metadata = {
        "filename": path.name,
        "file_type": ext.lstrip("."),
        "file_size_bytes": stat.st_size,
        "num_characters": len(text),
        "num_words": len(text.split()),
        "loaded_at": datetime.now(timezone.utc).isoformat(),
    }

    doc_id = generate_doc_id(str(path.resolve()))
    logger.info("Loaded document '%s' (%d chars)", path.name, len(text))
    return Document(doc_id=doc_id, source=str(path), text=text, metadata=metadata)


def load_documents(paths: List[str | Path]) -> List[Document]:
    """Load multiple documents, skipping (and logging) any that fail to parse."""
    documents: List[Document] = []
    for p in paths:
        try:
            documents.append(load_document(p))
        except DocumentLoadError as exc:
            logger.warning("Skipping document: %s", exc)
    return documents


def discover_documents(directory: str | Path) -> List[Path]:
    """Recursively find all supported document files under `directory`."""
    directory = Path(directory)
    if not directory.exists():
        return []
    return sorted(
        p for p in directory.rglob("*") if p.suffix.lower() in SUPPORTED_EXTENSIONS and p.is_file()
    )
